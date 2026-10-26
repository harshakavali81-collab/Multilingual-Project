"""API for the Multilingual Indian Language Knowledge Base Platform.

The app prefers LaBSE when the optional ML dependencies are installed. A small,
local hashing-vectorizer fallback keeps development, tests, and offline demos usable.
"""
from __future__ import annotations

import hashlib
import secrets
import json
import math
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Generator, Optional

from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field
from sqlalchemy import Column, DateTime, Integer, String, Text, create_engine, func, inspect, text
from sqlalchemy.orm import DeclarativeBase, Session, sessionmaker
from argon2 import PasswordHasher
from argon2.exceptions import VerificationError

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./local_kb.db")
_engine_options: dict = {"pool_pre_ping": True}
if DATABASE_URL.startswith("sqlite"):
    _engine_options["connect_args"] = {"check_same_thread": False}
engine = create_engine(DATABASE_URL, **_engine_options)
SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)


def _utcnow_naive() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


class Base(DeclarativeBase):
    pass


class DBDocument(Base):
    __tablename__ = "documents"
    id = Column(Integer, primary_key=True, index=True)
    title = Column(String(255), nullable=False)
    content = Column(Text, nullable=False)
    language = Column(String(20), nullable=False, default="English")
    category = Column(String(100), nullable=True, default="General")
    created_at = Column(DateTime, nullable=False, default=_utcnow_naive)
    owner_id = Column(Integer, nullable=True)


class ChatHistory(Base):
    __tablename__ = "chat_history"
    id = Column(Integer, primary_key=True, index=True)
    question = Column(Text, nullable=False)
    answer = Column(Text, nullable=False)
    language = Column(String(20), nullable=False, default="English")
    sources_json = Column(Text, nullable=False, default="[]")
    answer_mode = Column(String(30), nullable=False, default="extractive-fallback")
    created_at = Column(DateTime, nullable=False, default=_utcnow_naive)
    owner_id = Column(Integer, nullable=True)


class AuthUser(Base):
    __tablename__ = "auth_users"
    id = Column(Integer, primary_key=True)
    email = Column(String(254), unique=True, nullable=False, index=True)
    password_hash = Column(Text, nullable=False)
    role = Column(String(20), nullable=False, default="member")
    failed_logins = Column(Integer, nullable=False, default=0)
    locked_until = Column(DateTime, nullable=True)
    created_at = Column(DateTime, nullable=False, default=_utcnow_naive)


class LoginSession(Base):
    __tablename__ = "login_sessions"
    id = Column(Integer, primary_key=True)
    user_id = Column(Integer, nullable=False, index=True)
    token_hash = Column(String(64), unique=True, nullable=False, index=True)
    expires_at = Column(DateTime, nullable=False)
    created_at = Column(DateTime, nullable=False, default=_utcnow_naive)


Base.metadata.create_all(bind=engine)


def migrate_schema() -> None:
    """Idempotent, additive schema upgrade; back up the DB before deployment."""
    with engine.begin() as conn:
        for table_name in ("documents", "chat_history"):
            names = {col["name"] for col in inspect(conn).get_columns(table_name)}
            if "owner_id" not in names:
                conn.execute(text(f"ALTER TABLE {table_name} ADD COLUMN owner_id INTEGER"))
            conn.execute(text(f"CREATE INDEX IF NOT EXISTS ix_{table_name}_owner_id ON {table_name} (owner_id)"))


migrate_schema()
_hasher = PasswordHasher(time_cost=2, memory_cost=19456, parallelism=1)


SUPPORTED_LANGUAGES = {"en": "English", "hi": "Hindi", "te": "Telugu", "ta": "Tamil", "kn": "Kannada"}
LANGUAGE_CODES = {name.lower(): code for code, name in SUPPORTED_LANGUAGES.items()}
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", "10485760"))
MAX_DOCUMENT_CHARS = int(os.getenv("MAX_DOCUMENT_CHARS", "200000"))
EMBEDDING_DIMENSION = 384
_DOCUMENT_EMBEDDINGS: dict[int, list[float]] = {}
_encoder = None
_encoder_attempted = False


class QueryRequest(BaseModel):
    query: str = Field(min_length=1, max_length=4000)
    preferred_language: str = Field(default="auto", max_length=30)
    top_k: int = Field(default=3, ge=1, le=10)


class Credentials(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=1, max_length=128)


class NewUser(Credentials):
    role: str = Field(default="member")


class PasswordChange(BaseModel):
    current_password: str = Field(min_length=1)
    new_password: str = Field(min_length=12, max_length=128)


class UploadResponse(BaseModel):
    message: str
    document_id: int
    title: str
    detected_language: str


app = FastAPI(
    title="Multilingual Indian Language Knowledge Base Platform",
    description="Grounded document retrieval across English, Hindi, Telugu, Tamil, and Kannada. GROQ_API_KEY is optional.",
    version="1.0.0",
)

def auth_enabled() -> bool:
    return os.getenv("AUTH_REQUIRED", "false").lower() in ("1", "true", "yes")


def _normalize_email(email: str) -> str:
    email = email.strip().lower()
    if not 3 <= len(email) <= 254 or "@" not in email or email.startswith("@"):
        raise HTTPException(status_code=422, detail="A valid email address is required.")
    return email


def _principal(request: Request) -> Optional[int]:
    return getattr(request.state, "user_id", None) if auth_enabled() else None


def _user_scope(query, model, request: Request):
    if auth_enabled():
        return query.filter(model.owner_id == _principal(request))
    return query


def _public_user(user: AuthUser) -> dict:
    return {"id": user.id, "email": user.email, "role": user.role}


def bootstrap_admin() -> None:
    """First boot only: provision admin from secret environment variables."""
    if not auth_enabled():
        return
    with SessionLocal() as db:
        if db.query(AuthUser).count():
            return
        email = os.getenv("ADMIN_EMAIL", "").strip()
        password = os.getenv("ADMIN_PASSWORD", "")
        if not email or len(password) < 12:
            raise RuntimeError("AUTH_REQUIRED needs ADMIN_EMAIL and ADMIN_PASSWORD (12+ chars) on first boot.")
        db.add(AuthUser(email=_normalize_email(email),
                        password_hash=_hasher.hash(password), role="admin"))
        db.commit()


bootstrap_admin()


def _revoke_user_sessions(db: Session, user_id: int) -> None:
    db.query(LoginSession).filter(LoginSession.user_id == user_id).delete(synchronize_session=False)


def _require_admin(request: Request) -> None:
    if not auth_enabled() or getattr(request.state, "role", "") != "admin":
        raise HTTPException(status_code=403, detail="Administrator access required.")

@app.middleware("http")
async def optional_api_key(request: Request, call_next):
    """Require a shared key for data routes only when explicitly configured."""
    expected = os.getenv("APP_API_KEY", "").strip()
    public_paths = {"/", "/docs", "/openapi.json", "/redoc", "/health", "/auth/status", "/auth/login"}
    if request.method == "OPTIONS" or request.url.path in public_paths:
        return await call_next(request)
    if auth_enabled():
        header = request.headers.get("authorization", "")
        prefix = "Bearer "
        if not header.startswith(prefix) or len(header) <= len(prefix):
            return JSONResponse(status_code=401, content={"detail": "Please sign in."})
        token = header[len(prefix):]
        if len(token) > 256:
            return JSONResponse(status_code=401, content={"detail": "Invalid session."})
        token_hash = hashlib.sha256(token.encode()).hexdigest()
        with SessionLocal() as db:
            session = db.query(LoginSession).filter(
                LoginSession.token_hash == token_hash,
                LoginSession.expires_at > _utcnow_naive(),
            ).first()
            user = db.get(AuthUser, session.user_id) if session else None
            if not user:
                return JSONResponse(status_code=401, content={"detail": "Session expired. Please sign in."})
            request.state.user_id = user.id
            request.state.role = user.role
            request.state.session_id = session.id
        return await call_next(request)
    if expected:
        supplied = request.headers.get("x-api-key", "")
        if not secrets.compare_digest(supplied, expected):
            return JSONResponse(status_code=401, content={"detail": "Valid X-API-Key header required."})
    return await call_next(request)


_allowed_origins = [
    item.strip()
    for item in os.getenv("CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173,http://localhost:3000,http://127.0.0.1:3000").split(",")
    if item.strip()
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=_allowed_origins,
    allow_credentials=True,
    allow_methods=["GET", "POST", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type", "Authorization", "X-API-Key"],
)


def get_db() -> Generator[Session, None, None]:
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def detect_language(text: str) -> str:
    for char in text:
        codepoint = ord(char)
        if 0x0C00 <= codepoint <= 0x0C7F:
            return "te"
        if 0x0B80 <= codepoint <= 0x0BFF:
            return "ta"
        if 0x0C80 <= codepoint <= 0x0CFF:
            return "kn"
        if 0x0900 <= codepoint <= 0x097F:
            return "hi"
    try:
        from langdetect import detect
        guess = detect(text)
        return guess if guess in SUPPORTED_LANGUAGES else "en"
    except Exception:
        return "en"


def _hashing_embedding(text: str) -> list[float]:
    """Local lexical vector fallback. For cross-language semantics install LaBSE."""
    vector = [0.0] * EMBEDDING_DIMENSION
    tokens = re.findall(r"[^\W_]+", text.casefold(), flags=re.UNICODE)
    features: list[tuple[str, float]] = [(token, 1.0) for token in tokens]
    for token in tokens:
        padded = f"^{token}$"
        for size in (2, 3, 4):
            for index in range(max(0, len(padded) - size + 1)):
                features.append((padded[index:index + size], 0.25))
    for feature, weight in features:
        digest = hashlib.blake2b(feature.encode("utf-8"), digest_size=8).digest()
        number = int.from_bytes(digest, "big")
        vector[number % EMBEDDING_DIMENSION] += weight if number & 1 else -weight
    norm = math.sqrt(sum(value * value for value in vector))
    return [value / norm for value in vector] if norm else vector


def embed_text(text: str) -> list[float]:
    global _encoder, _encoder_attempted
    mode = os.getenv("EMBEDDING_BACKEND", "auto").strip().lower()
    if mode not in {"hashing", "local", "off"} and not _encoder_attempted:
        _encoder_attempted = True
        try:
            from sentence_transformers import SentenceTransformer
            _encoder = SentenceTransformer(os.getenv("EMBEDDING_MODEL", "sentence-transformers/LaBSE"))
        except Exception:
            _encoder = None
    if _encoder is not None and mode not in {"hashing", "local", "off"}:
        try:
            result = _encoder.encode(text, normalize_embeddings=True)
            return result.tolist() if hasattr(result, "tolist") else list(result)
        except Exception:
            _encoder = None
    return _hashing_embedding(text)


def _cosine_similarity(left: list[float], right: list[float]) -> float:
    if len(left) != len(right):
        return 0.0
    left_norm = math.sqrt(sum(value * value for value in left))
    right_norm = math.sqrt(sum(value * value for value in right))
    if not left_norm or not right_norm:
        return 0.0
    return sum(a * b for a, b in zip(left, right)) / (left_norm * right_norm)


def _document_embedding(document: DBDocument) -> list[float]:
    if document.id not in _DOCUMENT_EMBEDDINGS:
        _DOCUMENT_EMBEDDINGS[document.id] = embed_text(document.content)
    return _DOCUMENT_EMBEDDINGS[document.id]


def _language_name(preferred: str, query: str) -> str:
    candidate = (preferred or "auto").strip().lower()
    if candidate in {"auto", "detect", "automatic"}:
        candidate = detect_language(query)
    if candidate in SUPPORTED_LANGUAGES:
        return SUPPORTED_LANGUAGES[candidate]
    return SUPPORTED_LANGUAGES.get(LANGUAGE_CODES.get(candidate, ""), "English")


def _parse_document_file(filename: str, raw: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".txt", ".md", ".csv", ".pdf", ".docx"}:
        raise HTTPException(status_code=415, detail="Unsupported file type. Upload TXT, MD, CSV, PDF, or DOCX.")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"File exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MiB limit.")
    if not raw:
        raise HTTPException(status_code=422, detail="The uploaded file is empty.")
    if suffix == ".pdf":
        try:
            from io import BytesIO
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(raw))
            if len(reader.pages) > 250:
                raise HTTPException(status_code=413, detail="PDF exceeds the 250-page limit.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=422, detail="Could not read this PDF. Upload a valid text-based PDF.") from exc
        if not text:
            raise HTTPException(status_code=422, detail="No selectable text found. Scanned PDFs need OCR first.")
        return _validated_content(text)
    if suffix == ".docx":
        try:
            from io import BytesIO
            from docx import Document
            document = Document(BytesIO(raw))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            return _validated_content("\n".join(paragraphs))
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=422, detail="Could not read this DOCX document.") from exc
    try:
        return _validated_content(raw.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=422, detail="Text files must be encoded as UTF-8.") from exc


def _validated_content(value: str) -> str:
    """Reject empty or excessively large extracted text."""
    clean = value.strip()
    if not clean:
        raise HTTPException(status_code=422, detail="Document content cannot be empty.")
    if len(clean) > MAX_DOCUMENT_CHARS:
        raise HTTPException(status_code=413, detail="Extracted text exceeds the document character limit.")
    return clean


def _save_document(db: Session, title: str, content: str, category: str = "General", owner_id: Optional[int] = None) -> DBDocument:
    clean_title = (title or "Untitled document").strip()[:255] or "Untitled document"
    clean_content = _validated_content(content or "")
    document = DBDocument(
        title=clean_title,
        content=clean_content,
        language=SUPPORTED_LANGUAGES[detect_language(clean_content)],
        category=(category or "General").strip()[:100] or "General",
        owner_id=owner_id,
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    _DOCUMENT_EMBEDDINGS[document.id] = embed_text(clean_content)
    return document


def _serialize_history(item: ChatHistory) -> dict:
    try:
        sources = json.loads(item.sources_json or "[]")
    except (TypeError, json.JSONDecodeError):
        sources = []
    return {
        "id": item.id, "question": item.question, "answer": item.answer,
        "language": item.language, "answer_mode": item.answer_mode,
        "sources": sources, "created_at": item.created_at,
    }


def _local_answer(question: str, documents: list[DBDocument]) -> str:
    """Extract relevant sentences when no hosted LLM is configured."""
    candidates: list[tuple[int, str]] = []
    question_terms = set(re.findall(r"[^\W_]+", question.casefold(), flags=re.UNICODE))
    for document in documents:
        chunks = re.split(r"(?<=[.!?।])\s+|\n+", document.content.strip())
        for chunk in chunks:
            sentence = chunk.strip()
            if sentence:
                terms = set(re.findall(r"[^\W_]+", sentence.casefold(), flags=re.UNICODE))
                candidates.append((len(question_terms & terms), sentence))
    candidates.sort(key=lambda pair: pair[0], reverse=True)
    useful = [sentence for score, sentence in candidates if score > 0][:3]
    if useful:
        return " ".join(useful)
    if documents:
        return re.split(r"(?<=[.!?।])\s+|\n+", documents[0].content.strip())[0][:900]
    return "No documents are available yet. Upload a document, then ask again."


def _generate_answer(question: str, context: str, language: str, documents: list[DBDocument]) -> tuple[str, str, Optional[str]]:
    api_key = os.getenv("GROQ_API_KEY", "").strip()
    if not api_key:
        return _local_answer(question, documents), "extractive-fallback", "GROQ_API_KEY is not configured; returned an extractive answer."
    try:
        from groq import Groq
        completion = Groq(api_key=api_key).chat.completions.create(
            model=os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile"),
            messages=[
                {"role": "system", "content": (
                    "Answer only from the supplied context. Do not invent facts. "
                    "If context is insufficient, say so. Keep concise. Answer in " + language + "."
                )},
                {"role": "user", "content": "Context:\n" + context + "\n\nQuestion:\n" + question},
            ],
            temperature=0.1,
        )
        answer = (completion.choices[0].message.content or "").strip()
        if not answer:
            raise RuntimeError("Empty response from answer provider")
        return answer, "groq-llm", None
    except Exception:
        return _local_answer(question, documents), "extractive-fallback", "The LLM provider was unavailable; returned an extractive answer."



@app.get("/auth/status")
def auth_status() -> dict:
    return {"required": auth_enabled(), "password_login": auth_enabled()}


@app.post("/auth/login")
def login(credentials: Credentials, db: Session = Depends(get_db)) -> dict:
    if not auth_enabled():
        raise HTTPException(status_code=404, detail="Password login not enabled.")
    email = _normalize_email(credentials.email)
    user = db.query(AuthUser).filter(AuthUser.email == email).first()
    if not user:
        # Ensure an unknown account still incurs password-hash work.
        try:
            _hasher.verify(_hasher.hash("dummy-unknown-password"), credentials.password)
        except VerificationError:
            pass
        raise HTTPException(status_code=401, detail="Invalid credentials.")
    now = _utcnow_naive()
    if user.locked_until and user.locked_until > now:
        raise HTTPException(status_code=429, detail="Account temporarily locked. Try again later.")
    try:
        _hasher.verify(user.password_hash, credentials.password)
    except VerificationError:
        user.failed_logins = (user.failed_logins or 0) + 1
        if user.failed_logins >= 5:
            user.failed_logins = 0
            user.locked_until = now + timedelta(minutes=15)
        db.commit()
        raise HTTPException(status_code=401, detail="Invalid credentials.")
    user.failed_logins = 0
    user.locked_until = None
    token = secrets.token_urlsafe(32)
    db.add(LoginSession(
        user_id=user.id, token_hash=hashlib.sha256(token.encode()).hexdigest(),
        expires_at=now + timedelta(hours=8),
    ))
    db.commit()
    return {"access_token": token, "token_type": "bearer", "expires_in": 28800,
            "user": _public_user(user)}


@app.get("/auth/me")
def current_user(request: Request, db: Session = Depends(get_db)) -> dict:
    if not auth_enabled():
        raise HTTPException(status_code=404, detail="Password login not enabled.")
    user = db.get(AuthUser, _principal(request))
    if not user:
        raise HTTPException(status_code=401, detail="Session expired.")
    return _public_user(user)


@app.post("/auth/logout")
def logout(request: Request, db: Session = Depends(get_db)) -> dict:
    if not auth_enabled():
        raise HTTPException(status_code=404, detail="Password login not enabled.")
    session = db.get(LoginSession, request.state.session_id)
    if session:
        db.delete(session)
        db.commit()
    return {"message": "Signed out."}


@app.post("/auth/change-password")
def change_password(payload: PasswordChange, request: Request, db: Session = Depends(get_db)) -> dict:
    if not auth_enabled():
        raise HTTPException(status_code=404, detail="Password login not enabled.")
    user = db.get(AuthUser, _principal(request))
    try:
        _hasher.verify(user.password_hash, payload.current_password)
    except VerificationError:
        raise HTTPException(status_code=401, detail="Incorrect current password.")
    user.password_hash = _hasher.hash(payload.new_password)
    _revoke_user_sessions(db, user.id)
    db.commit()
    return {"message": "Password changed. Please sign in again."}


@app.get("/admin/users")
def list_users(request: Request, db: Session = Depends(get_db)) -> list[dict]:
    _require_admin(request)
    return [_public_user(user) for user in db.query(AuthUser).order_by(AuthUser.id).all()]


@app.post("/admin/users", status_code=201)
def create_user(payload: NewUser, request: Request, db: Session = Depends(get_db)) -> dict:
    _require_admin(request)
    if payload.role not in ("member", "admin"):
        raise HTTPException(status_code=422, detail="Role must be member or admin.")
    if len(payload.password) < 12:
        raise HTTPException(status_code=422, detail="Password requires at least 12 characters.")
    email = _normalize_email(payload.email)
    if db.query(AuthUser).filter(AuthUser.email == email).first():
        raise HTTPException(status_code=409, detail="Account already exists.")
    user = AuthUser(email=email, password_hash=_hasher.hash(payload.password), role=payload.role)
    db.add(user)
    db.commit()
    db.refresh(user)
    return _public_user(user)


@app.get("/")
def root() -> dict:
    return {"name": app.title, "version": app.version, "docs": "/docs", "health": "/health"}


@app.get("/health")
def system_health() -> dict:
    return {
        "status": "operational",
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "groq_configured": bool(os.getenv("GROQ_API_KEY", "").strip()),
        "authentication_required": auth_enabled(),
        "embedding_backend": os.getenv("EMBEDDING_BACKEND", "auto"),
    }


@app.get("/articles", response_model=None)
def read_all_articles(request: Request, db: Session = Depends(get_db)) -> list[DBDocument]:
    return _user_scope(db.query(DBDocument), DBDocument, request).order_by(DBDocument.created_at.desc(), DBDocument.id.desc()).all()


@app.get("/documents/{document_id}", response_model=None)
def read_document(document_id: int, request: Request, db: Session = Depends(get_db)) -> DBDocument:
    document = _user_scope(db.query(DBDocument), DBDocument, request).filter(DBDocument.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")
    return document


@app.post("/upload", response_model=UploadResponse, status_code=201)
def ingest_document(
    request: Request,
    title: str = Form(..., min_length=1, max_length=255),
    content: str = Form(..., min_length=1),
    category: str = Form("General", max_length=100),
    db: Session = Depends(get_db),
) -> UploadResponse:
    document = _save_document(db, title, content, category, _principal(request))
    return UploadResponse(message="Document saved and indexed.", document_id=document.id, title=document.title, detected_language=document.language)


@app.post("/upload-file", response_model=UploadResponse, status_code=201)
async def upload_file(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)) -> UploadResponse:
    filename = Path(file.filename or "uploaded.txt").name
    raw = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"File exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MiB limit.")
    content = _parse_document_file(filename, raw)
    document = _save_document(db, filename, content, "Uploaded", _principal(request))
    return UploadResponse(message="File uploaded, parsed, saved, and indexed.", document_id=document.id, title=document.title, detected_language=document.language)


@app.delete("/documents/{document_id}")
def delete_document(document_id: int, request: Request, db: Session = Depends(get_db)) -> dict:
    document = _user_scope(db.query(DBDocument), DBDocument, request).filter(DBDocument.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")
    db.delete(document)
    db.commit()
    _DOCUMENT_EMBEDDINGS.pop(document_id, None)
    return {"message": "Document deleted.", "document_id": document_id}


@app.post("/ask")
def cross_language_query(payload: QueryRequest, request: Request, db: Session = Depends(get_db)) -> dict:
    question = payload.query.strip()
    if not question:
        raise HTTPException(status_code=422, detail="Question cannot be empty.")
    documents = _user_scope(db.query(DBDocument), DBDocument, request).all()
    language = _language_name(payload.preferred_language, question)
    if not documents:
        answer = "The knowledge base is empty. Upload a document before asking a question."
        result = {
            "answer": answer, "sources": [],
            "detected_query_language": detect_language(question),
            "answer_language": language, "answer_mode": "empty-knowledge-base", "notice": None,
        }
    else:
        query_embedding = embed_text(question)
        ranked = sorted(
            ((_cosine_similarity(query_embedding, _document_embedding(doc)), doc) for doc in documents),
            key=lambda item: item[0],
            reverse=True,
        )
        selected = [doc for score, doc in ranked[:payload.top_k] if score >= float(os.getenv("MIN_RETRIEVAL_SCORE", "0.04"))]
        if not selected:
            result = {
                "answer": "I could not find relevant information in the uploaded documents. Try rephrasing or adding a relevant document.",
                "sources": [], "detected_query_language": detect_language(question),
                "answer_language": language, "answer_mode": "no-relevant-context", "notice": None,
            }
        else:
            context = "\n\n".join(
                "Source: " + doc.title + "\nLanguage: " + doc.language + "\nContent:\n" + doc.content[:12000]
                for doc in selected
            )
            answer, mode, notice = _generate_answer(question, context, language, selected)
            result = {
                "answer": answer,
                "sources": [{"id": doc.id, "title": doc.title, "language": doc.language, "category": doc.category} for doc in selected],
                "detected_query_language": detect_language(question),
                "answer_language": language, "answer_mode": mode, "notice": notice,
            }
    history = ChatHistory(
        question=question, answer=result["answer"], language=result["answer_language"],
        sources_json=json.dumps(result["sources"], ensure_ascii=False),
        answer_mode=result["answer_mode"],
        owner_id=_principal(request),
    )
    db.add(history)
    db.commit()
    return result


@app.get("/history")
def read_history(request: Request, limit: int = 100, db: Session = Depends(get_db)) -> list[dict]:
    bounded_limit = max(1, min(limit, 500))
    items = _user_scope(db.query(ChatHistory), ChatHistory, request).order_by(ChatHistory.created_at.desc(), ChatHistory.id.desc()).limit(bounded_limit).all()
    return [_serialize_history(item) for item in items]


@app.delete("/history")
def clear_history(request: Request, db: Session = Depends(get_db)) -> dict:
    deleted = _user_scope(db.query(ChatHistory), ChatHistory, request).delete(synchronize_session=False)
    db.commit()
    return {"message": "Chat history cleared.", "deleted_count": deleted}


@app.get("/search")
def search_documents(
    request: Request,
    q: str = Query(min_length=1, max_length=4000),
    limit: int = Query(default=10, ge=1, le=50),
    db: Session = Depends(get_db),
) -> dict:
    """Search document metadata and text without exposing entire documents."""
    query = q.strip()
    if not query:
        raise HTTPException(status_code=422, detail="Search query cannot be empty.")
    vector = embed_text(query)
    rows = _user_scope(db.query(DBDocument), DBDocument, request).all()
    ranked = sorted(
        ((_cosine_similarity(vector, _document_embedding(item)), item) for item in rows),
        key=lambda pair: pair[0], reverse=True,
    )
    matches = [
        {
            "id": item.id,
            "title": item.title,
            "language": item.language,
            "category": item.category,
            "score": round(max(0.0, score), 4),
            "preview": item.content[:220],
        }
        for score, item in ranked if score > 0
    ]
    return {"query": query, "total": len(matches), "results": matches[:limit]}


@app.get("/analytics")
def get_analytics(request: Request, days: int = Query(default=14, ge=1, le=90), db: Session = Depends(get_db)) -> dict:
    """Return real document/question activity for a chart-ready dashboard."""
    now = datetime.now(timezone.utc).date()
    cutoff = datetime.combine(now - timedelta(days=days - 1), datetime.min.time())
    document_rows = _user_scope(
        db.query(func.date(DBDocument.created_at), func.count(DBDocument.id)), DBDocument, request
    ).filter(DBDocument.created_at >= cutoff).group_by(func.date(DBDocument.created_at)).all()
    question_rows = _user_scope(
        db.query(func.date(ChatHistory.created_at), func.count(ChatHistory.id)), ChatHistory, request
    ).filter(ChatHistory.created_at >= cutoff).group_by(func.date(ChatHistory.created_at)).all()
    doc_map = {str(day): count for day, count in document_rows}
    question_map = {str(day): count for day, count in question_rows}
    daily = []
    for offset in range(days - 1, -1, -1):
        day = (now - timedelta(days=offset)).isoformat()
        daily.append({"date": day, "documents": doc_map.get(day, 0), "questions": question_map.get(day, 0)})
    categories = _user_scope(db.query(DBDocument.category, func.count(DBDocument.id)), DBDocument, request).group_by(DBDocument.category).all()
    modes = _user_scope(db.query(ChatHistory.answer_mode, func.count(ChatHistory.id)), ChatHistory, request).group_by(ChatHistory.answer_mode).all()
    languages = _user_scope(db.query(DBDocument.language, func.count(DBDocument.id)), DBDocument, request).group_by(DBDocument.language).all()
    return {
        "days": days,
        "daily": daily,
        "language_counts": {lang: count for lang, count in languages},
        "category_counts": {category or "General": count for category, count in categories},
        "answer_modes": {mode: count for mode, count in modes},
    }


@app.get("/stats")
def get_stats(request: Request, db: Session = Depends(get_db)) -> dict:
    document_count = _user_scope(db.query(func.count(DBDocument.id)), DBDocument, request).scalar() or 0
    question_count = _user_scope(db.query(func.count(ChatHistory.id)), ChatHistory, request).scalar() or 0
    rows = _user_scope(db.query(DBDocument.language, func.count(DBDocument.id)), DBDocument, request).group_by(DBDocument.language).all()
    return {
        "document_count": document_count,
        "question_count": question_count,
        "language_counts": {language: count for language, count in rows},
        "supported_languages": SUPPORTED_LANGUAGES,
    }
